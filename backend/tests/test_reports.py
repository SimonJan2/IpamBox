"""V11 Reports workspace — /reports/summary, per-section CSV, multi-sheet
XLSX, the email slice, site scoping, truncation, empty DB, RBAC.

The reuse contract is the test: every section number is asserted against
the same services/pages the report claims to summarize
(dashboard_stats, stamp_rack_stats, review predicates, due_where)."""
import io
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.redis import get_redis
from app.core.security import hash_password
from app.models.cabling import DeviceInterface
from app.models.certificate import Certificate
from app.models.device import Device
from app.models.ip_address import IPAddress, IPStatus
from app.models.monitoring import (
    MonitorKind,
    MonitorState,
    MonitorTarget,
    NotificationChannel,
    ChannelKind,
)
from app.models.prefix import Prefix, PrefixStatus
from app.models.rack import Rack, RackGroup
from app.models.scan_job import ScanJob, ScanStatus
from app.models.site import Site
from app.models.user import User, UserRole
from app.models.vrf import VRF
from app.services import ipam, reports

PASSWORD = "reports-test-pw1"


@pytest.fixture
async def auth_on():
    settings = get_settings()
    settings.ipambox_allow_insecure = False
    r = get_redis()
    try:
        yield
    finally:
        settings.ipambox_allow_insecure = True
        for pattern in ("ipam:session:*", "ipam:loginfails:*", "ipam:lockout:*"):
            async for key in r.scan_iter(pattern):
                await r.delete(key)
        await r.aclose()


async def _mkuser(session: AsyncSession, username: str, role: UserRole) -> User:
    u = User(username=username, password_hash=hash_password(PASSWORD), role=role)
    session.add(u)
    await session.commit()
    return u


async def _login(client: AsyncClient, username: str) -> None:
    r = await client.post(
        "/api/v1/auth/login", json={"username": username, "password": PASSWORD}
    )
    assert r.status_code == 200, r.text


@pytest.fixture
async def seeded(session: AsyncSession) -> dict:
    """Two sites worth of estate: prefixes in both families, addresses in
    every status, racks+capacity, certs on both sides of the warn window,
    scan history, monitors, and one of each flag kind."""
    today = date.today()
    hq = Site(name="HQ", slug="hq")
    br = Site(name="Branch", slug="branch")
    session.add_all([hq, br])
    await session.flush()
    vrf = await session.scalar(select(VRF).where(VRF.name == "Global"))
    vrf.site_id = hq.id

    p_hq = Prefix(prefix="10.71.0.0/24", vrf_id=vrf.id, site_id=hq.id,
                  status=PrefixStatus.ACTIVE)
    p_br = Prefix(prefix="192.168.9.0/24", vrf_id=vrf.id, site_id=br.id,
                  status=PrefixStatus.ACTIVE)
    p_cont = Prefix(prefix="10.99.0.0/16", vrf_id=vrf.id, site_id=hq.id,
                    status=PrefixStatus.CONTAINER)
    p_v6 = Prefix(prefix="fd00:71::/64", vrf_id=vrf.id, site_id=br.id,
                  status=PrefixStatus.ACTIVE)
    session.add_all([p_hq, p_br, p_cont, p_v6])
    await session.flush()

    rg = RackGroup(name="DC-1", site_id=hq.id)
    session.add(rg)
    await session.flush()
    r1 = Rack(name="R1", site_id=hq.id, group_id=rg.id, height_u=42)
    r2 = Rack(name="R2", site_id=br.id, height_u=24)
    session.add_all([r1, r2])
    await session.flush()

    d1 = Device(name="core-sw", site_id=hq.id, rack_id=r1.id, u_position=1,
                u_height=2, category="network", watts=350, weight_kg=8.5)
    d2 = Device(name="srv1", site_id=hq.id, rack_id=r1.id, u_position=10,
                u_height=1, category="server", watts=500, weight_kg=12.0)
    d3 = Device(name="bench-pc", site_id=br.id, category="endpoint")
    session.add_all([d1, d2, d3])
    await session.flush()

    def ip(prefix, addr, status, device=None, mac=None, custom=None):
        return IPAddress(
            address=addr,
            address_int=int(__import__("ipaddress").ip_address(addr)),
            prefix_id=prefix.id,
            vrf_id=prefix.vrf_id,
            status=status,
            device_id=device,
            mac_address=mac,
            custom_fields=custom,
        )

    dup_mac = "aa:bb:cc:00:00:01"
    addrs = [
        ip(p_hq, "10.71.0.1", IPStatus.ACTIVE, device=d1.id, mac=dup_mac),
        ip(p_hq, "10.71.0.2", IPStatus.DISCOVERED),
        ip(p_hq, "10.71.0.3", IPStatus.OFFLINE),
        ip(p_hq, "10.71.0.4", IPStatus.ACTIVE, mac=dup_mac),
        ip(p_hq, "10.71.0.5", IPStatus.ACTIVE,
           custom={"mac_mismatch": {"was": "aa:bb:cc:00:00:09",
                                    "seen": "aa:bb:cc:00:00:10",
                                    "at": datetime.now(timezone.utc).isoformat()}}),
        ip(p_br, "192.168.9.1", IPStatus.RESERVED),
    ]
    session.add_all(addrs)
    await session.flush()
    session.add(
        DeviceInterface(
            device_id=d1.id, name="Gi0/1", kind="rj45",
            validation={"cable_mismatch": {"reason": "neighbor"}},
        )
    )
    session.add_all([
        Certificate(cert_name="expired.example", server_name="expired.example",
                    expires_on=today - timedelta(days=5)),
        Certificate(cert_name="soon.example", server_name="soon.example",
                    expires_on=today + timedelta(days=10)),
        Certificate(cert_name="fine.example", server_name="fine.example",
                    expires_on=today + timedelta(days=200)),
    ])
    session.add_all([
        ScanJob(cidr="10.71.0.0/24", status=ScanStatus.COMPLETED,
                hosts_discovered=40, hosts_new=4,
                finished_at=datetime.now(timezone.utc), duration_seconds=9.5),
        ScanJob(cidr="10.71.0.0/24", status=ScanStatus.FAILED,
                finished_at=datetime.now(timezone.utc), duration_seconds=1.0),
        ScanJob(cidr="192.168.9.0/24", status=ScanStatus.COMPLETED,
                hosts_discovered=12, hosts_new=2,
                finished_at=datetime.now(timezone.utc), duration_seconds=4.0),
    ])
    session.add_all([
        MonitorTarget(device_id=d1.id, kind=MonitorKind.PING,
                      state=MonitorState.UP),
        MonitorTarget(address_id=addrs[0].id, kind=MonitorKind.HTTP,
                      state=MonitorState.DOWN),
        MonitorTarget(device_id=d3.id, kind=MonitorKind.TCP,
                      state=MonitorState.UNKNOWN, enabled=False),
    ])
    await session.commit()
    return {"hq": hq, "br": br, "p_hq": p_hq, "p_br": p_br,
            "r1": r1, "r2": r2, "rg": rg, "d1": d1, "d2": d2, "d3": d3}


def _sec(report: dict, key: str) -> dict:
    return next(s for s in report["sections"] if s["key"] == key)


async def _summary(client: AsyncClient, **params) -> dict:
    r = await client.get("/api/v1/reports/summary", params=params)
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------------- summary


async def test_summary_matches_service_aggregates(
    client: AsyncClient, session: AsyncSession, seeded
):
    report = await _summary(client)
    stats = await ipam.dashboard_stats(session)

    assert [s["key"] for s in report["sections"]] == [
        "sites", "status", "utilization", "capacity", "certs", "scans",
        "flags", "devices", "monitors",  # audits gated off — no v9 table
    ]
    assert report["site"] is None

    # headline strip == dashboard_stats
    m = {x["label"]: x["value"] for x in report["metrics"]}
    assert m["Sites"] == stats["sites_total"] == 2
    assert m["VRFs"] == stats["vrfs_total"] == 1
    assert m["Prefixes"] == stats["prefixes_total"] == 4
    assert m["Addresses"] == f"{stats['ips_used']} used / {stats['ips_total']} usable"
    assert m["Utilization"] == f"{stats['utilization_pct']}%"
    assert m["Racks"] == stats["racks_total"] == 2
    assert m["Rack U"] == f"{stats['rack_u_used']} / {stats['rack_u_total']}"

    # sites: per-site fill math matches the dashboard's aggregate
    sites = {r[0]: r for r in _sec(report, "sites")["rows"]}
    assert sites["HQ"][1:] == [2, 5, 254, 2.0, 2]
    assert sites["Branch"][1:] == [2, 1, 254, 0.4, 1]

    # status == dashboard's by_status buckets
    status = {r[0]: r[1] for r in _sec(report, "status")["rows"]}
    assert status["active"] == stats["devices_active"] == 3
    assert status["discovered"] == stats["devices_discovered"] == 1
    assert status["offline"] == stats["devices_offline"] == 1
    assert status["reserved"] == stats["devices_reserved"] == 1

    # capacity == stamp_rack_stats rolled up per group
    cap = {r[0]: r for r in _sec(report, "capacity")["rows"]}
    assert cap["DC-1"][2:] == [1, 3, 39, 42, 7.1, 850, 20.5]
    assert cap["(ungrouped)"][2:] == [1, 0, 24, 24, 0.0, None, None]

    # certs: expired + inside-window only
    certs = _sec(report, "certs")
    assert certs["total_rows"] == stats["certs_expiring_30d"] == 2
    assert [r[6] for r in certs["rows"]] == ["expired", "expiring"]
    cm = {x["label"]: x["value"] for x in certs["metrics"]}
    assert cm["Certificates"] == stats["certificates_total"] == 3
    assert cm["Expired"] == 1

    # scans: status counts + hosts seen
    sm = {x["label"]: x["value"] for x in _sec(report, "scans")["metrics"]}
    assert sm["Jobs"] == stats["scans_total"] == 3
    assert sm["Completed"] == 2 and sm["Failed"] == 1
    assert sm["Hosts seen"] == 52
    assert [r[2] for r in _sec(report, "scans")["rows"]] == [
        "completed", "failed", "completed",
    ]

    # flags == review predicates (detection state)
    flags = {r[0]: r[2] for r in _sec(report, "flags")["rows"]}
    assert flags["mac_mismatch"] == stats["mac_mismatches"] == 1
    assert flags["cable_mismatch"] == stats["cable_mismatches"] == 1
    assert flags["dup_mac"] == 1
    assert flags["discovered_unconfirmed"] == 1

    # devices: category / health / placement
    dev = _sec(report, "devices")
    by_dim = {(r[0], r[1]): r[2] for r in dev["rows"]}
    assert by_dim[("category", "network")] == 1
    assert by_dim[("health", "active")] == 1  # d1 owns an ACTIVE address
    assert by_dim[("health", "unmonitored")] == 2
    assert by_dim[("placement", "unracked")] == 1
    dm = {x["label"]: x["value"] for x in dev["metrics"]}
    assert dm["Devices"] == 3 and dm["Unracked"] == 1

    # monitors: state/kind/enabled rollup + due count
    mon = _sec(report, "monitors")
    mm = {x["label"]: x["value"] for x in mon["metrics"]}
    assert mm["Up"] == 1 and mm["Down"] == 1 and mm["Unknown"] == 1
    assert mm["Due now"] == 2  # enabled + never checked
    mdim = {(r[0], r[1]): r[2] for r in mon["rows"]}
    assert mdim[("kind", "ping")] == 1 and mdim[("kind", "http")] == 1
    assert mdim[("kind", "tcp")] == 1
    assert mdim[("enabled", "no")] == 1


async def test_utilization_rows_sort_full_and_empty(
    client: AsyncClient, session: AsyncSession, seeded
):
    report = await _summary(client)
    u = _sec(report, "utilization")
    fullest = [r for r in u["rows"] if r[0] == "fullest"]
    emptiest = [r for r in u["rows"] if r[0] == "emptiest"]
    # /24s scored (containers + v6 excluded); fullest is the HQ prefix
    assert fullest[0][1] == "10.71.0.0/24" and fullest[0][7] == 2.0
    assert emptiest[0][1] == "192.168.9.0/24" and emptiest[0][7] == 0.4
    # prefix_stats_dict parity on the scored rows
    for r in fullest + emptiest:
        assert round(100 * r[5] / r[6], 1) == r[7]


async def test_site_scoping_narrows_everything(
    client: AsyncClient, session: AsyncSession, seeded
):
    hq, br = seeded["hq"], seeded["br"]
    scoped = await _summary(client, site_id=hq.id)
    assert scoped["site"]["name"] == "HQ"
    m = {x["label"]: x["value"] for x in scoped["metrics"]}
    assert m["Sites"] == 1 and m["Prefixes"] == 2 and m["Devices"] == 2
    assert m["Addresses"] == "5 used / 254 usable"
    assert m["Rack U"] == "3 / 42"

    sites = _sec(scoped, "sites")["rows"]
    assert len(sites) == 1 and sites[0][0] == "HQ"
    flags = {r[0]: r[2] for r in _sec(scoped, "flags")["rows"]}
    assert flags == {"mac_mismatch": 1, "cable_mismatch": 1,
                     "dup_mac": 1, "discovered_unconfirmed": 1}
    # branch's reserved address falls out of the status section
    status = {r[0]: r[1] for r in _sec(scoped, "status")["rows"]}
    assert status["reserved"] == 0 and status["active"] == 3
    mon = {x["label"]: x["value"] for x in _sec(scoped, "monitors")["metrics"]}
    assert mon["Up"] == 1 and mon["Unknown"] == 0  # d3 lives on Branch

    scoped_br = await _summary(client, site_id=br.id)
    assert {r[0] for r in _sec(scoped_br, "sites")["rows"]} == {"Branch"}
    br_flags = {r[0]: r[2] for r in _sec(scoped_br, "flags")["rows"]}
    assert all(v == 0 for v in br_flags.values())

    assert (await client.get(
        "/api/v1/reports/summary", params={"site_id": 9999}
    )).status_code == 404


async def test_empty_db_zeroes_sections(client: AsyncClient):
    report = await _summary(client)
    keys = [s["key"] for s in report["sections"]]
    # tables exist → sections present but empty, not absent/errored
    assert "sites" in keys and "monitors" in keys and "audits" not in keys
    status = {r[0]: r[1] for r in _sec(report, "status")["rows"]}
    assert all(v == 0 for v in status.values())
    assert all(r[2] == 0 for r in _sec(report, "flags")["rows"])
    assert all(r[2] == 0 for r in _sec(report, "devices")["rows"])
    assert all(r[2] == 0 for r in _sec(report, "monitors")["rows"])
    for key in ("sites", "utilization", "capacity", "certs", "scans"):
        assert _sec(report, key)["rows"] == [], key
    m = {x["label"]: x["value"] for x in report["metrics"]}
    assert m["Utilization"] == "0.0%" and m["Sites"] == 0


async def test_truncation_marker(client: AsyncClient, session: AsyncSession, monkeypatch):
    monkeypatch.setattr(reports, "SECTION_ROWS", 2)
    vrf = await session.scalar(select(VRF).where(VRF.name == "Global"))
    session.add_all([
        Certificate(cert_name=f"c{i}", expires_on=date.today() + timedelta(days=1))
        for i in range(3)
    ])
    session.add(Prefix(prefix="10.80.0.0/24", vrf_id=vrf.id,
                       status=PrefixStatus.ACTIVE))
    await session.commit()
    report = await _summary(client)
    certs = _sec(report, "certs")
    assert certs["truncated"] is True
    assert certs["total_rows"] == 3 and len(certs["rows"]) == 2


# --------------------------------------------------------------------- xlsx


async def test_xlsx_workbook_opens_with_ordered_sheets(
    client: AsyncClient, seeded
):
    summary = await _summary(client)
    r = await client.get("/api/v1/reports/export.xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    assert "ipambox-report-" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == [s["title"] for s in summary["sections"]]
    for s in summary["sections"]:
        ws = wb[s["title"]]
        rows = list(ws.iter_rows(values_only=True))
        assert list(rows[0]) == s["columns"]
        assert len(rows) - 1 == len(s["rows"])  # capped rows, honest marker
    wb.close()


# ---------------------------------------------------------------------- csv


async def test_section_csv_bom_headers_and_parity(
    client: AsyncClient, seeded
):
    summary = await _summary(client)
    for s in summary["sections"]:
        r = await client.get(f"/api/v1/reports/{s['key']}.csv")
        assert r.status_code == 200, s["key"]
        body = r.content.decode("utf-8")
        assert body.startswith("\ufeff")  # Hebrew-in-Excel BOM convention
        lines = [ln for ln in body.lstrip("\ufeff").splitlines() if ln]
        assert lines[0] == ",".join(s["columns"])
        assert len(lines) - 1 == len(s["rows"])

    assert (await client.get("/api/v1/reports/nope.csv")).status_code == 404


async def test_csv_site_scoping(client: AsyncClient, seeded):
    r = await client.get(
        "/api/v1/reports/sites.csv", params={"site_id": seeded["hq"].id}
    )
    assert r.status_code == 200
    body = r.text
    assert "HQ" in body and "Branch" not in body
    assert f"report-sites-{seeded['hq'].slug}-" in r.headers["content-disposition"]


# -------------------------------------------------------------------- email


async def test_email_fans_out_through_channels(
    client: AsyncClient, session: AsyncSession, seeded, monkeypatch
):
    sent = []

    async def fake_emit(event_type, summary, payload=None):
        sent.append((event_type, summary, payload))
        return 1

    monkeypatch.setattr("app.services.notify.emit", fake_emit)
    session.add(
        NotificationChannel(
            name="ops", kind=ChannelKind.WEBHOOK,
            config={"url": "https://example.invalid/hook"}, enabled=True,
        )
    )
    await session.commit()

    r = await client.post("/api/v1/reports/email")
    assert r.status_code == 200, r.text
    assert r.json() == {"channels": 1, "event": "report.requested"}
    event, summary, payload = sent[0]
    assert event == "report.requested"
    assert "all sites" in summary and "Flags: 4" in summary
    assert payload["url"] == "/reports"
    assert payload["sections"]["devices"] > 0

    # scoped send names the site
    await client.post(
        "/api/v1/reports/email", params={"site_id": seeded["hq"].id}
    )
    assert "HQ" in sent[-1][1]
    assert "site_id=1" in sent[-1][2]["url"] or "site_id=" in sent[-1][2]["url"]


# --------------------------------------------------------------------- rbac


async def test_reports_rbac(
    client: AsyncClient, session: AsyncSession, auth_on, monkeypatch
):
    async def fake_emit(event_type, summary, payload=None):
        return 0

    monkeypatch.setattr("app.services.notify.emit", fake_emit)
    await _mkuser(session, "v", UserRole.VIEWER)
    await _mkuser(session, "op", UserRole.OPERATOR)
    await _login(client, "v")

    # viewers read every report surface
    assert (await client.get("/api/v1/reports/summary")).status_code == 200
    assert (await client.get("/api/v1/reports/export.xlsx")).status_code == 200
    assert (await client.get("/api/v1/reports/sites.csv")).status_code == 200
    # ...but sending a report is a write action
    assert (await client.post("/api/v1/reports/email")).status_code == 403

    await _login(client, "op")
    assert (await client.post("/api/v1/reports/email")).status_code == 200


async def test_reports_write_no_changelog(
    client: AsyncClient, session: AsyncSession, seeded, monkeypatch
):
    """Reports are derived — reads and email writes never hit the audit log."""
    from app.models.change_log import ChangeLog

    async def fake_emit(event_type, summary, payload=None):
        return 0

    monkeypatch.setattr("app.services.notify.emit", fake_emit)

    before = int(await session.scalar(select(func.count(ChangeLog.id))) or 0)
    await _summary(client)
    await client.get("/api/v1/reports/export.xlsx")
    await client.post("/api/v1/reports/email")
    after = int(await session.scalar(select(func.count(ChangeLog.id))) or 0)
    assert after == before
