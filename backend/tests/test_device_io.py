"""V5B: /devices export.csv + export.xlsx + smart /devices/import.

The contract: what the page shows is what exports; what exports
re-imports cleanly. Covers filter parity, column order, BOM, FK names,
xlsx read-back, header auto-map (canonical/alias/NetBox/Hebrew), match
precedence, honest diffs, placement conflicts, carrier two-pass, IP
linking, dry-run/force semantics, round-trip, and RBAC/changelog.
"""

import io
import json

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.redis import get_redis
from app.core.security import hash_password
from app.models.user import User, UserRole
from app.services.device_io import DEVICE_EXPORT_COLUMNS

PASSWORD = "io-test-pw1"


# ------------------------------------------------------------------ helpers


@pytest.fixture
async def auth_on():
    settings = get_settings()
    settings.ipambox_allow_insecure = False
    cookie_secure = settings.ipambox_cookie_secure
    settings.ipambox_cookie_secure = False
    r = get_redis()
    try:
        yield
    finally:
        settings.ipambox_allow_insecure = True
        settings.ipambox_cookie_secure = cookie_secure
        for pattern in ("ipam:session:*", "ipam:loginfails:*", "ipam:lockout:*"):
            async for key in r.scan_iter(pattern):
                await r.delete(key)
        await r.aclose()


async def _mkuser(session: AsyncSession, username: str, role: UserRole) -> User:
    u = User(username=username, password_hash=hash_password(PASSWORD), role=role)
    session.add(u)
    await session.commit()
    return u


async def _login(client: AsyncClient, username: str) -> None:
    r = await client.post(
        "/api/v1/auth/login", json={"username": username, "password": PASSWORD}
    )
    assert r.status_code == 200, r.text


async def _site(client: AsyncClient, name="DC-1") -> dict:
    r = await client.post("/api/v1/sites", json={"name": name})
    assert r.status_code == 201, r.text
    return r.json()


async def _group(client: AsyncClient, name="Row A") -> dict:
    r = await client.post("/api/v1/rack-groups", json={"name": name})
    assert r.status_code == 201, r.text
    return r.json()


async def _rack(client: AsyncClient, **kw) -> dict:
    body = {"name": "R1", "height_u": 12, "width": 19, **kw}
    r = await client.post("/api/v1/racks", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def _device(client: AsyncClient, **kw) -> dict:
    body = {"name": "srv", **kw}
    r = await client.post("/api/v1/devices", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def _racked(client: AsyncClient, rack_id: int, **kw) -> dict:
    body = {"name": "srv", "u_position": 1, "u_height": 1, "face": "front", **kw}
    r = await client.post(f"/api/v1/racks/{rack_id}/devices", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def _prefix(client: AsyncClient, cidr="10.70.0.0/24") -> int:
    vrfs = (await client.get("/api/v1/vrfs")).json()
    vrf = (vrfs["items"] if isinstance(vrfs, dict) else vrfs)[0]
    r = await client.post(
        "/api/v1/prefixes", json={"prefix": cidr, "vrf_id": vrf["id"]}
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


async def _ip(client: AsyncClient, prefix_id: int, addr: str, **kw) -> dict:
    r = await client.post(
        "/api/v1/addresses", json={"address": addr, "prefix_id": prefix_id, **kw}
    )
    assert r.status_code == 201, r.text
    return r.json()


def _csv(text: str) -> bytes:
    return text.encode("utf-8")


def _xlsx(headers: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _import(
    client: AsyncClient, payload: bytes, filename="devices.csv", **params
) -> dict:
    r = await client.post(
        "/api/v1/devices/import",
        content=payload,
        params={"filename": filename, **params},
    )
    assert r.status_code == 200, r.text
    return r.json()


def _actions(resp: dict) -> dict:
    return {r["row"]: r["action"] for r in resp["rows"]}


# ------------------------------------------------------------------- export


async def test_export_csv_column_order_bom_and_names(client: AsyncClient):
    site = await _site(client)
    group = await _group(client)
    rack = await _rack(client, name="R-EXP", site_id=site["id"],
                       group_id=group["id"])
    pid = await _prefix(client)
    d = await _device(
        client, name="exp1", rack_id=rack["id"], site_id=site["id"],
        u_position=2, u_height=2,
        manufacturer="Dell", model="R650", serial_number="SN-E1",
        device_type="server", category="compute", notes="n1",
    )
    await _ip(client, pid, "10.70.0.10", device_id=d["id"])
    await _ip(client, pid, "10.70.0.11", device_id=d["id"])

    r = await client.get("/api/v1/devices/export.csv")
    assert r.status_code == 200
    body = r.content
    assert body.startswith(b"\xef\xbb\xbf")  # utf-8 BOM
    lines = body.decode("utf-8-sig").splitlines()
    header = lines[0].split(",")
    assert header == DEVICE_EXPORT_COLUMNS
    row = lines[1].split(",")
    rec = dict(zip(header, row))
    assert rec["site"] == "DC-1" and rec["rack"] == "R-EXP"
    assert rec["rack_group"] == "Row A"
    assert rec["ips"] == "10.70.0.10 10.70.0.11"
    assert rec["ip_count"] == "2"
    assert rec["id"] == str(d["id"])
    assert "devices-filtered" not in r.headers["content-disposition"]


async def test_export_filter_parity_with_list(client: AsyncClient):
    rack = await _rack(client, name="R-F")
    await _racked(client, rack["id"], name="in-rack", u_position=1)
    await _device(client, name="free", manufacturer="HP")
    await _device(client, name="dellbox", manufacturer="Dell")

    params = {"manufacturer": "dell"}
    listed = (
        await client.get("/api/v1/devices", params=params)
    ).json()["items"]
    r = await client.get("/api/v1/devices/export.csv", params=params)
    assert "devices-filtered.csv" in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").splitlines()
    assert {l.split(",")[0] for l in lines[1:]} == {
        str(d["id"]) for d in listed
    } == {str((await client.get("/api/v1/devices",
                                params={"q": "dellbox"})).json()["items"][0]["id"])}

    # wiring facet: same post-aggregate path as the list
    params = {"wiring": "uncabled"}
    listed = (await client.get("/api/v1/devices", params=params)).json()["items"]
    r = await client.get("/api/v1/devices/export.csv", params=params)
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) - 1 == len(listed)

    # health facet
    params = {"health": "unmonitored"}
    listed = (await client.get("/api/v1/devices", params=params)).json()["items"]
    r = await client.get("/api/v1/devices/export.csv", params=params)
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) - 1 == len(listed) == 3


async def test_export_columns_whitelist(client: AsyncClient):
    await _device(client, name="lean")
    r = await client.get(
        "/api/v1/devices/export.csv",
        params={"columns": "name,device_type,id,name"},
    )
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0].split(",") == ["name", "device_type", "id"]
    r = await client.get(
        "/api/v1/devices/export.csv", params={"columns": "name,bogus"}
    )
    assert r.status_code == 422


async def test_export_xlsx_matches_csv(client: AsyncClient):
    rack = await _rack(client, name="RX")
    await _racked(client, rack["id"], name="x1", u_position=3)
    await _device(client, name="x2")

    csv_r = await client.get("/api/v1/devices/export.csv")
    xlsx_r = await client.get("/api/v1/devices/export.xlsx")
    assert xlsx_r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_r.content))
    assert wb.sheetnames == ["devices"]
    ws = wb["devices"]
    x_rows = [[("" if c is None else str(c)) for c in row]
              for row in ws.iter_rows(values_only=True)]
    c_lines = csv_r.content.decode("utf-8-sig").splitlines()
    c_rows = [l.split(",") for l in c_lines]
    assert x_rows[0] == c_rows[0] == DEVICE_EXPORT_COLUMNS
    assert x_rows[1:] == c_rows[1:]


# ------------------------------------------------------------- header map


async def test_import_detect_canonical_alias_hebrew(client: AsyncClient):
    # canonical headers
    payload = _csv("id,name,rack,u_position\n9,a,R,3\n")
    resp = await _import(client, payload, detect=True)
    cols = {c["header"]: c["field"] for c in resp["columns"]}
    assert cols == {"id": "id", "name": "name", "rack": "rack",
                    "u_position": "u_position"}

    # English aliases
    payload = _csv("hostname,serial,rack u,height,mac\nh,S1,5,2,\n")
    resp = await _import(client, payload, detect=True)
    cols = {c["header"]: c["field"] for c in resp["columns"]}
    assert cols["hostname"] == "name"
    assert cols["serial"] == "serial_number"
    assert cols["rack u"] == "u_position"
    assert cols["height"] == "u_height"
    assert cols["mac"] == "mac_address"

    # Hebrew headers
    payload = _csv("שם,אתר,ארון,מיקום u,הערות\nשרת1,אתר-א,R9,4,ה\n")
    resp = await _import(client, payload, detect=True)
    cols = {c["header"]: c["field"] for c in resp["columns"]}
    assert cols["שם"] == "name" and cols["אתר"] == "site"
    assert cols["ארון"] == "rack" and cols["מיקום u"] == "u_position"
    assert cols["הערות"] == "notes"

    # NetBox shape: device_role + device_type => model, not our type
    payload = _csv("name,device_role,device_type,site,rack,position\n"
                   "n1,server,dell-r650,S1,R2,5\n")
    resp = await _import(client, payload, detect=True)
    cols = {c["header"]: c["field"] for c in resp["columns"]}
    assert cols["device_role"] == "device_type"
    assert cols["device_type"] == "model"
    assert cols["position"] == "u_position"

    # mapping override wins over auto-map; unknown header rejected
    resp = await _import(
        client, payload,
        mapping=json.dumps({"device_type": "device_type"}),
        detect=True,
    )
    cols = {c["header"]: c["field"] for c in resp["columns"]}
    assert cols["device_type"] == "device_type"
    r = await client.post(
        "/api/v1/devices/import",
        content=payload,
        params={"filename": "d.csv", "detect": True,
                "mapping": json.dumps({"nosuch": "name"})},
    )
    assert r.status_code == 422


# --------------------------------------------------------------- dry-run


async def test_import_dry_run_writes_nothing(client: AsyncClient):
    resp = await _import(client, _csv("name,device_type\nn1,server\n"))
    assert resp["counts"] == {"create": 1, "update": 0, "skip": 0, "error": 0}
    assert resp["committed"] is False
    assert (await client.get("/api/v1/devices")).json()["total"] == 0

    # even commit-mode rows with errors write nothing without force
    resp = await _import(
        client, _csv("name,rack,u_position\ngood,NOPE,1\n"), dry_run=False
    )
    assert resp["counts"]["error"] == 1 and resp["committed"] is False
    assert (await client.get("/api/v1/devices")).json()["total"] == 0


async def test_import_commit_create_and_changelog(client: AsyncClient):
    rack = await _rack(client, name="RC")
    pid = await _prefix(client)
    await _ip(client, pid, "10.70.0.50")
    payload = _csv(
        "name,rack,u_position,u_height,face,ips,notes\n"
        f"n1,RC,2,2,rear,10.70.0.50,hi\n"
        f"n2,RC,5,1,,10.70.0.99,\n"
    )
    resp = await _import(client, payload, dry_run=False)
    assert resp["committed"] is True
    assert resp["counts"] == {"create": 2, "update": 0, "skip": 0, "error": 0}
    detail = {r["row"]: r["detail"] for r in resp["rows"]}
    assert "unknown ip '10.70.0.99'" in detail[3]

    devices = (await client.get("/api/v1/devices")).json()["items"]
    by_name = {d["name"]: d for d in devices}
    assert by_name["n1"]["u_position"] == 2
    assert by_name["n1"]["face"] == "rear"
    assert by_name["n1"]["source"] == "import"
    got = (
        await client.get(f"/api/v1/devices/{by_name['n1']['id']}")
    ).json()
    assert [i["address"] for i in got["ips"]] == ["10.70.0.50"]

    log = (
        await client.get("/api/v1/changelog", params={"object_type": "Device"})
    ).json()["items"]
    assert any(
        e["action"] == "create" and e["object_repr"] == "n1" for e in log
    )


async def test_import_force_commits_ok_rows(client: AsyncClient):
    payload = _csv(
        "name,rack,u_position\nok1,,,\nbad1,NOPE,2\n"
    )
    resp = await _import(client, payload, dry_run=False, force=True)
    assert resp["committed"] is True
    assert resp["counts"]["create"] == 1 and resp["counts"]["error"] == 1
    names = {
        d["name"]
        for d in (await client.get("/api/v1/devices")).json()["items"]
    }
    assert names == {"ok1"}


# ---------------------------------------------------------- match/update


async def test_import_match_precedence_and_skip(client: AsyncClient):
    d1 = await _device(client, name="same", serial_number="SN-A")
    d2 = await _device(client, name="other", serial_number="SN-B",
                       mac_address="AA:BB:CC:DD:EE:01")
    payload = _csv(
        "id,name,serial_number,mac_address\n"
        f"{d1['id']},same,SN-B,AA:BB:CC:DD:EE:01\n"   # id wins over serial/mac
        f",same,SN-B,AA:BB:CC:DD:EE:01\n"            # serial wins over mac
        f",x,,AA:BB:CC:DD:EE:01\n"                   # mac match
        ",same,,\n"                                 # name match
    )
    resp = await _import(client, payload)  # on_match=skip default
    rows = {r["row"]: r for r in resp["rows"]}
    assert rows[2]["action"] == "skip" and "matched by id" in rows[2]["detail"]
    assert "matched by serial_number" in rows[3]["detail"]
    assert "matched by mac_address" in rows[4]["detail"]
    assert "matched by name" in rows[5]["detail"]

    # ambiguous name -> error, not a coin flip
    await _device(client, name="same")  # second device named "same"
    resp = await _import(client, _csv("name\nsame\n"))
    assert resp["rows"][0]["action"] == "error"
    assert "ambiguous" in resp["rows"][0]["detail"]


async def test_import_update_diffs_are_honest(client: AsyncClient):
    rack = await _rack(client, name="RU")
    d = await _device(client, name="upd", manufacturer="Dell",
                      serial_number="SN-U")
    payload = _csv(
        "serial_number,name,manufacturer,model\n"
        "SN-U,upd-renamed,Dell,R750\n"
    )
    resp = await _import(client, payload, on_match="update")
    row = resp["rows"][0]
    assert row["action"] == "update"
    assert row["diff"] == {
        "name": ["upd", "upd-renamed"], "model": [None, "R750"]
    }
    # commit applies the diff, no more
    resp = await _import(client, payload, on_match="update", dry_run=False)
    assert resp["committed"] is True
    got = (await client.get(f"/api/v1/devices/{d['id']}")).json()
    assert got["name"] == "upd-renamed" and got["model"] == "R750"
    assert got["manufacturer"] == "Dell"

    # identical file -> skip
    payload2 = _csv(
        "serial_number,name,manufacturer,model\n"
        "SN-U,upd-renamed,Dell,R750\n"
    )
    resp = await _import(client, payload2, on_match="update")
    assert resp["rows"][0]["action"] == "skip"
    assert "identical" in resp["rows"][0]["detail"]


# --------------------------------------------------------------- placement


async def test_import_placement_conflict_names_occupant(client: AsyncClient):
    rack = await _rack(client, name="RP")
    await _racked(client, rack["id"], name="blocker", u_position=4)
    payload = _csv("name,rack,u_position,face\nnew1,RP,4,front\n")
    resp = await _import(client, payload)
    row = resp["rows"][0]
    assert row["action"] == "error" and "blocker" in row["detail"]

    # conflicts against earlier batch rows are also caught
    payload = _csv(
        "name,rack,u_position,face\nnew1,RP,7,front\nnew2,RP,7,front\n"
    )
    resp = await _import(client, payload)
    rows = {r["row"]: r for r in resp["rows"]}
    assert rows[2]["action"] == "create"
    assert rows[3]["action"] == "error" and "new1" in rows[3]["detail"]


async def test_import_unracked_on_missing(client: AsyncClient):
    payload = _csv("name,rack,u_position\nu1,GHOST,3\n")
    resp = await _import(client, payload)
    assert resp["rows"][0]["action"] == "error"
    resp = await _import(client, payload, unracked_on_missing=True,
                         dry_run=False)
    assert resp["committed"] is True
    d = (await client.get("/api/v1/devices")).json()["items"][0]
    assert d["rack_id"] is None

    # a broken rack_group qualifier voids the placement claim — same
    # salvage as a missing rack (the rack cell isn't trusted either)
    payload = _csv("name,rack,rack_group,u_position\nu2,G1,NOGROUP,4\n")
    resp = await _import(client, payload)
    assert resp["rows"][0]["action"] == "error"
    assert "rack_group" in resp["rows"][0]["detail"]
    resp = await _import(client, payload, unracked_on_missing=True,
                         dry_run=False)
    assert resp["committed"] is True
    assert "unracked" in resp["rows"][0]["detail"]
    devs = {
        x["name"]: x
        for x in (await client.get("/api/v1/devices")).json()["items"]
    }
    assert devs["u2"]["rack_id"] is None

    # pass B: a child row with a broken group still mounts — the carrier's
    # own placement decides, the rack cell is just voided
    rack = await _rack(client, name="RCY")
    await _racked(client, rack["id"], name="tray2", u_position=5,
                  slot_layout="halves")
    payload = _csv("name,rack,rack_group,carrier,slot\nkid3,,NOGROUP,tray2,0\n")
    resp = await _import(client, payload, unracked_on_missing=True,
                         dry_run=False)
    assert resp["committed"] is True, resp["rows"]
    devs = {
        x["name"]: x
        for x in (await client.get("/api/v1/devices")).json()["items"]
    }
    assert devs["kid3"]["rack_id"] == rack["id"] and devs["kid3"]["slot"] == 0


# ---------------------------------------------------------------- carriers


async def test_import_carrier_two_pass(client: AsyncClient):
    rack = await _rack(client, name="RCX")
    # children listed BEFORE their carrier in the file — two-pass handles it
    payload = _csv(
        "name,rack,u_position,slot_layout,carrier,slot\n"
        "kid1,RCX,,,tray,0\n"
        "tray,RCX,8,halves,,\n"
        "kid2,RCX,,,tray,1\n"
    )
    resp = await _import(client, payload)
    rows = {r["row"]: r for r in resp["rows"]}
    assert rows[2]["action"] == "create"
    assert rows[3]["action"] == "create"
    assert rows[4]["action"] == "create"

    resp = await _import(client, payload, dry_run=False)
    assert resp["committed"] is True, resp
    devices = (await client.get("/api/v1/devices")).json()["items"]
    by_name = {d["name"]: d for d in devices}
    tray = by_name["tray"]
    for kid in ("kid1", "kid2"):
        assert by_name[kid]["rack_id"] == rack["id"]
        assert by_name[kid]["u_position"] == 8  # inherited from carrier
        assert by_name[kid]["carrier_id"] == tray["id"]
    assert by_name["kid1"]["slot"] == 0 and by_name["kid2"]["slot"] == 1

    # missing carrier -> error row
    resp = await _import(client, _csv("name,carrier,slot\nk,NOPE,0\n"))
    assert resp["rows"][0]["action"] == "error"


async def test_import_replaces_unracked_carrier_and_children(
    client: AsyncClient,
):
    """Devices-first-then-racks flow: an unracked carrier + child already
    in inventory must be placeable by a later import. on_match=update
    re-racks the carrier and mounts the child; on_match=skip leaves both
    untouched (the child reports skip, not a misleading not-found)."""
    rack = await _rack(client, name="RL")
    await _device(client, name="trayU", slot_layout="halves")
    await _device(client, name="kidU")
    payload = _csv(
        "name,rack,u_position,slot_layout,carrier,slot\n"
        "trayU,RL,7,halves,,\n"
        "kidU,RL,,,trayU,1\n"
    )

    resp = await _import(client, payload, on_match="skip")
    rows = {r["row"]: r for r in resp["rows"]}
    assert rows[2]["action"] == "skip"
    assert rows[3]["action"] == "skip" and "not racked" in rows[3]["detail"]

    resp = await _import(client, payload, on_match="update", dry_run=False)
    assert resp["committed"] is True, resp["rows"]
    rows = {r["row"]: r for r in resp["rows"]}
    assert rows[2]["action"] == "update"
    assert rows[3]["action"] == "update"
    devs = {
        d["name"]: d
        for d in (await client.get("/api/v1/devices")).json()["items"]
    }
    tray = devs["trayU"]
    assert tray["rack_id"] == rack["id"] and tray["u_position"] == 7
    assert devs["kidU"]["carrier_id"] == tray["id"]
    assert devs["kidU"]["slot"] == 1
    assert devs["kidU"]["rack_id"] == rack["id"]


# --------------------------------------------------------------- ip links


async def test_import_ip_link_and_warn(client: AsyncClient):
    pid = await _prefix(client)
    ip = await _ip(client, pid, "10.70.0.77")
    resp = await _import(
        client, _csv("name,ips\nwithip,10.70.0.77 10.70.0.78\n"),
        dry_run=False,
    )
    assert resp["committed"] is True
    row = resp["rows"][0]
    assert "unknown ip '10.70.0.78'" in row["detail"]
    d = (await client.get("/api/v1/devices")).json()["items"][0]
    got = (await client.get(f"/api/v1/devices/{d['id']}")).json()
    assert [i["id"] for i in got["ips"]] == [ip["id"]]
    # no address was created
    addrs = (await client.get("/api/v1/addresses")).json()
    assert len(addrs) == 1


# --------------------------------------------------------------- xlsx in


async def test_import_xlsx_first_sheet(client: AsyncClient):
    payload = _xlsx(
        ["name", "manufacturer", "watts"],
        [["xl1", "HP", 300], ["xl2", "Dell", 400]],
    )
    resp = await _import(client, payload, filename="devices.xlsx",
                         dry_run=False)
    assert resp["committed"] is True and resp["counts"]["create"] == 2
    names = {
        d["name"] for d in (await client.get("/api/v1/devices")).json()["items"]
    }
    assert names == {"xl1", "xl2"}


# -------------------------------------------------------------- round-trip


async def test_round_trip_export_import_all_identical(client: AsyncClient):
    site = await _site(client)
    rack = await _rack(client, name="RRT", site_id=site["id"])
    pid = await _prefix(client)
    d1 = await _racked(
        client, rack["id"], name="rt1", u_position=2, u_height=2,
        manufacturer="Dell", model="R650", serial_number="SN-RT1",
        device_type="server", watts=350, notes="rt",
    )
    await _ip(client, pid, "10.70.0.90", device_id=d1["id"])
    await _device(client, name="rt2", manufacturer="HP", category="net")

    r = await client.get("/api/v1/devices/export.csv")
    payload = r.content  # BOM'd utf-8 — load_upload strips it
    resp = await _import(client, payload, on_match="update")
    assert resp["counts"]["error"] == 0
    assert resp["counts"]["create"] == 0
    assert resp["counts"]["update"] == 0
    assert resp["counts"]["skip"] == 2
    assert all("identical" in row["detail"] for row in resp["rows"])

    # and through xlsx too
    r = await client.get("/api/v1/devices/export.xlsx")
    resp = await _import(client, r.content, filename="devices.xlsx",
                         on_match="update")
    assert resp["counts"]["skip"] == 2 and resp["counts"]["error"] == 0


# ---------------------------------------------------------------------- RBAC


async def test_import_rbac(
    client: AsyncClient, session: AsyncSession, auth_on
):
    await _mkuser(session, "v", UserRole.VIEWER)
    await _mkuser(session, "c", UserRole.CONTRIBUTOR)
    payload = _csv("name\nrbac1\n")

    await _login(client, "v")
    r = await client.post(
        "/api/v1/devices/import", content=payload,
        params={"filename": "d.csv"},
    )
    assert r.status_code == 403
    # viewers can still export
    assert (await client.get("/api/v1/devices/export.csv")).status_code == 200
    assert (await client.get("/api/v1/devices/export.xlsx")).status_code == 200

    await _login(client, "c")
    r = await client.post(
        "/api/v1/devices/import", content=payload,
        params={"filename": "d.csv", "dry_run": "false"},
    )
    assert r.status_code == 200 and r.json()["committed"] is True
